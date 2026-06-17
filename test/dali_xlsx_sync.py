#!/usr/bin/env python3
"""Sync DALI device identities between the hardware and `Switch wiring.xlsx`.

The DALI **serial number** (memory bank 0) is a stable per-device id; the short
**address** changes whenever you (re-)provision.  This tool uses the serial as
the anchor so a mapping survives re-addressing.  It works on the "Lights" sheet,
which has `Device` (gateway Modbus slave id) and `Address` (DALI short address)
columns; it adds `Serial` and `GTIN` columns if they aren't there yet.

Two modes:

  populate  — scan each gateway, and for every device write its GTIN + Serial
              into the matching (Device, Address) row.  Run this while the
              addresses are known-good, to anchor each light to a physical
              device.

  remap     — after a re-provision changed addresses: scan each gateway for the
              current serial -> address map, then rewrite each row's Address by
              matching its recorded Serial.  Re-establishes all the mappings.

    .venv/bin/python test/dali_xlsx_sync.py populate
    .venv/bin/python test/dali_xlsx_sync.py remap --dry-run
    .venv/bin/python test/dali_xlsx_sync.py remap

A timestamped .bak copy is written before saving.  NOTE: openpyxl preserves cell
values/formulas but may drop charts/images — eyeball the file after the first run.
"""
import argparse
import glob
import os
import re
import shutil
import sys
import time

import serial
import openpyxl

LIGHTS_SHEET = "Lights"
DEFAULT_XLSX = os.path.join(os.path.dirname(__file__), "..", "Switch wiring.xlsx")


# ---- controller CLI ------------------------------------------------------
def cli(ser, cmd, wait):
    ser.reset_input_buffer()
    ser.write((cmd + "\n").encode())
    ser.flush()
    time.sleep(wait)
    return [l for l in ser.read(8000).decode("utf-8", "replace").splitlines()
            if l.strip() and not l.startswith("#") and not l.startswith("{")]


def identify(ser, slave, addr):
    """Return (gtin, serial) hex strings, or None for absent/collision."""
    line = (cli(ser, f"identify {slave} {addr}", 8.0) or [""])[-1]
    m = re.search(r"gtin=([0-9A-Fa-f]+) serial=([0-9A-Fa-f]+)", line)
    if m:
        return m.group(1).upper(), m.group(2).upper()
    if "collision" in line:
        print(f"    addr {addr}: COLLISION (resolve before syncing)")
    return None


def present_addrs(ser, slave):
    """Trigger a fresh gateway scan and return the present short addresses."""
    cli(ser, f"set_holding {slave} 0x1010 0x0001", 0.5)
    time.sleep(3.0)
    body = "\n".join(cli(ser, f"get_input {slave} 0x0100 4", 0.6))
    regs = [int(m.group(2), 16)
            for m in re.finditer(r'^\s*(\d+)\s+\d+\s+0x([0-9A-Fa-f]{4})\s*$', body, re.M)]
    mask = 0
    for i, v in enumerate(regs[:4]):
        mask |= v << (16 * i)
    return [a for a in range(64) if mask & (1 << a)]


# ---- spreadsheet helpers -------------------------------------------------
def col_index(ws, name):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v and str(v).strip().lower() == name.lower():
            return c
    return None


def ensure_col(ws, name):
    c = col_index(ws, name)
    if c:
        return c
    c = ws.max_column + 1
    ws.cell(1, c).value = name
    return c


def as_int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def main():
    ap = argparse.ArgumentParser(description="Sync DALI serials <-> Switch wiring.xlsx")
    ap.add_argument("mode", choices=["populate", "remap"])
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--port", default=None)
    ap.add_argument("--dry-run", action="store_true", help="scan + report, don't write the file")
    cfg = ap.parse_args()

    wb = openpyxl.load_workbook(cfg.xlsx)
    if LIGHTS_SHEET not in wb.sheetnames:
        sys.exit(f"no '{LIGHTS_SHEET}' sheet in {cfg.xlsx}")
    ws = wb[LIGHTS_SHEET]
    c_dev = col_index(ws, "Device")
    c_addr = col_index(ws, "Address")
    if not c_dev or not c_addr:
        sys.exit("Lights sheet needs 'Device' and 'Address' columns")
    c_ser = ensure_col(ws, "Serial")
    c_gtin = ensure_col(ws, "GTIN")

    port = cfg.port or sorted(glob.glob("/dev/tty.usbmodem*") + glob.glob("/dev/ttyACM*"))[0]
    ser = serial.Serial(port, 115200, timeout=0.2)
    time.sleep(0.5)
    ser.reset_input_buffer()
    cli(ser, "log warn", 0.4)

    # Collect the rows we care about.
    rows = []
    for r in range(2, ws.max_row + 1):
        dev = as_int(ws.cell(r, c_dev).value)
        if dev is None:
            continue
        addr = as_int(ws.cell(r, c_addr).value)
        s = ws.cell(r, c_ser).value
        rows.append((r, dev, addr, str(s).strip().upper() if s else None))

    # Scan every gateway referenced in the sheet: (dev, addr) -> (gtin, serial).
    gateways = sorted({dev for (_, dev, _, _) in rows})
    found = {}
    for dev in gateways:
        print(f"scanning gateway {dev}...")
        for a in present_addrs(ser, dev):
            res = identify(ser, dev, a)
            if res:
                found[(dev, a)] = res
                print(f"    addr {a}: gtin {res[0]} serial {res[1]}")
    ser.close()
    by_serial = {(dev, ser_no): a for (dev, a), (_, ser_no) in found.items()}

    changed = 0
    if cfg.mode == "populate":
        for (r, dev, addr, _) in rows:
            if addr is None:
                continue
            res = found.get((dev, addr))
            if not res:
                continue
            gtin, ser_no = res
            light = ws.cell(r, 1).value
            old = str(ws.cell(r, c_ser).value or "").upper()
            ws.cell(r, c_ser).value = ser_no
            ws.cell(r, c_gtin).value = gtin
            tag = "" if old == ser_no else "  (updated)"
            print(f"  row {r:3d} {str(light):30.30} dev {dev} addr {addr} -> serial {ser_no}{tag}")
            changed += 1
        print(f"\npopulated {changed} row(s) with serial + GTIN")
    else:  # remap
        missing = 0
        for (r, dev, addr, ser_no) in rows:
            if not ser_no:
                continue
            new_addr = by_serial.get((dev, ser_no))
            light = ws.cell(r, 1).value
            if new_addr is None:
                missing += 1
                print(f"  row {r:3d} {str(light):30.30} serial {ser_no}: NOT FOUND on gw {dev} (left at {addr})")
                continue
            if addr != new_addr:
                ws.cell(r, c_addr).value = new_addr
                print(f"  row {r:3d} {str(light):30.30} serial {ser_no}: addr {addr} -> {new_addr}")
                changed += 1
        print(f"\nremapped {changed} address(es); {missing} recorded serial(s) not found on the bus")

    if cfg.dry_run:
        print("dry-run: spreadsheet NOT written")
        return
    if changed == 0:
        print("no changes; spreadsheet not written")
        return
    bak = f"{cfg.xlsx}.{time.strftime('%Y%m%d-%H%M%S')}.bak"
    shutil.copy2(cfg.xlsx, bak)
    wb.save(cfg.xlsx)
    print(f"saved {cfg.xlsx}  (backup: {bak})")
    print("NOTE: openpyxl may not preserve charts/images — check the file.")


if __name__ == "__main__":
    try:
        main()
    except (IndexError, KeyboardInterrupt) as e:
        sys.exit(str(e) or "interrupted")
