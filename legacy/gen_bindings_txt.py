#!/usr/bin/env python3
"""gen_bindings_txt.py — generate a Lua-era /bindings.txt from Switch wiring.xlsx.

Reconstructed from the (removed) bridge/sync_bindings.py: it reuses that tool's
proven spreadsheet parsing (Switch Positions remap + Lights bindings) verbatim
and swaps only the output stage — from `replace_mappings` JSON over USB CDC to
the positional bindings.txt the old Lua firmware (src/main.lua, also removed)
read from /bindings.txt at boot.

bindings.txt format (consumed by main.lua's bindings.read()):
  * one line per (fixture*8 + button) index, 0 .. (24*8 - 1)
  * each line is a comma-separated list of `device*256 + address` targets,
    where device = Modbus slave and address = DALI short address (0..63)
  * a blank line means that button has no bindings
  e.g. line index 1 (fixture 0, button 1) "780,791" = {slave 3 addr 12, slave 3 addr 23}

Historical workflow: generate bindings.txt here, copy it onto the device's
littlefs as /bindings.txt (the old sync.py uploaded files over serial), reboot.

NOTE: the CURRENT controller firmware is the C firmware and does NOT read
/bindings.txt — this targets the historical Lua firmware. It lives in legacy/
for reference and one-off regeneration and is wired into nothing (no CMake).

Dependencies: openpyxl. Run with the repo venv, e.g.:
  .venv/bin/python legacy/gen_bindings_txt.py            # writes ./bindings.txt
  .venv/bin/python legacy/gen_bindings_txt.py --stdout   # print instead
"""

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl


MAX_FIXTURES        = 24
MAX_BUTTONS         = 8       # 7 physical + 1 spare
MAX_TARGETS_PER_BTN = 8

@dataclass
class Target:
    slave: int
    addr: int

    def to_json(self) -> dict:
        return {"slave": self.slave, "addr": self.addr}


# ── Spreadsheet parsing ──────────────────────────────────────────────────────

def load_remap(wb) -> tuple[dict, list[str]]:
    """Read Switch Positions tab into a remap dict.

    Returns ({(fixture_id, fix_name): [phys_at_logical_0, ...]}, errors).
    Rows whose `Patch Panel Location` (col B) is not an integer are skipped
    silently — they're commentary rows.
    """
    ws = wb["Switch Positions"]
    remap: dict[tuple[int, str], list[int]] = {}
    errors: list[str] = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            continue  # header
        name, fix_id, *order_cells = row[:7]
        if not isinstance(fix_id, int):
            continue  # commentary / blank
        if not (0 <= fix_id < MAX_FIXTURES):
            errors.append(f"Switch Positions row {row_idx}: fixture {fix_id} "
                          f"out of range 0..{MAX_FIXTURES - 1}")
            continue

        order: list[int] = []
        for col_idx, cell in enumerate(order_cells):
            if cell is None or cell == "":
                continue
            try:
                phys = int(cell)
            except (TypeError, ValueError):
                errors.append(f"Switch Positions row {row_idx} col {col_idx + 3}: "
                              f"non-integer button {cell!r}")
                continue
            if not (0 <= phys < MAX_BUTTONS):
                errors.append(f"Switch Positions row {row_idx} col {col_idx + 3}: "
                              f"physical button {phys} out of range")
                continue
            order.append(phys)

        key = (fix_id, name)
        if key in remap:
            errors.append(f"Switch Positions row {row_idx}: duplicate "
                          f"(fixture {fix_id}, name {name!r})")
            continue
        remap[key] = order

    # Disjoint check for shared fixture IDs.
    by_fix: dict[int, dict[str, list[int]]] = defaultdict(dict)
    for (fix_id, name), order in remap.items():
        by_fix[fix_id][name] = order
    for fix_id, plates in by_fix.items():
        if len(plates) <= 1:
            continue
        seen: dict[int, str] = {}
        for name, order in plates.items():
            for phys in order:
                if phys in seen and seen[phys] != name:
                    errors.append(f"Fixture {fix_id}: physical button {phys} "
                                  f"claimed by both {seen[phys]!r} and {name!r}")
                seen[phys] = name

    return remap, errors


def load_bindings(wb, remap: dict) -> tuple[dict, list[str]]:
    """Read Lights tab and invert into per-(fixture, physical_button) bindings.

    Returns ({(fixture_id, phys_button): [Target, ...]}, errors).

    The fixture ID for each binding is looked up by name in `remap` — the
    spreadsheet only stores the name and logical button per binding. Plate
    names are unique even when two plates share a fixture ID. """
    ws = wb["Lights"]
    bindings: dict[tuple[int, int], list[Target]] = defaultdict(list)
    errors: list[str] = []

    # name → fixture_id from Switch Positions. Plate names are globally
    # unique, so name disambiguates between plates that share a fixture ID.
    name_to_fix_id: dict[str, int] = {}
    name_collisions: list[str] = []
    for (fix_id, name) in remap:
        if name in name_to_fix_id and name_to_fix_id[name] != fix_id:
            name_collisions.append(
                f"Switch Positions: name {name!r} maps to both fixture "
                f"{name_to_fix_id[name]} and {fix_id}")
        else:
            name_to_fix_id[name] = fix_id
    errors.extend(name_collisions)

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            continue  # header
        # Columns: A Light, B Location, C Device, D Address,
        #          E Fix Name, F Btn1, G Fix Name 2, H Btn2,
        #          I Fix name 3, J Btn3, K Note
        if len(row) < 10:
            continue
        light_name = row[0]
        device     = row[2]
        address    = row[3]

        # Skip rows where the light isn't fully wired up — placeholders with
        # missing Device or Address can't drive a binding either way, so
        # silently ignore rather than blocking the sync.
        if device is None or address is None:
            continue

        if not isinstance(device, int) or not (1 <= device <= 247):
            errors.append(f"Lights row {row_idx} ({light_name!r}): "
                          f"invalid Device {device!r}")
            continue
        slave = device
        if not isinstance(address, int) or not (0 <= address < 64):
            errors.append(f"Lights row {row_idx} ({light_name!r}): "
                          f"invalid Address {address!r} (must be 0..63)")
            continue

        for triplet_idx, base in enumerate((4, 6, 8)):
            fix_name = row[base]
            log_btn  = row[base + 1]
            if fix_name is None and log_btn is None:
                continue
            # Partial pair (name without button or vice versa) means the row
            # is mid-edit — silently skip. A fully-set but invalid pair
            # (e.g. unknown fix_name) still errors below.
            if fix_name is None or not isinstance(log_btn, int):
                continue
            fix_id = name_to_fix_id.get(fix_name)
            if fix_id is None:
                errors.append(f"Lights row {row_idx} ({light_name!r}) "
                              f"binding {triplet_idx + 1}: no Switch Positions "
                              f"entry for plate {fix_name!r}")
                continue
            order = remap[(fix_id, fix_name)]
            if not (0 <= log_btn < len(order)):
                errors.append(f"Lights row {row_idx} ({light_name!r}) "
                              f"binding {triplet_idx + 1}: logical button "
                              f"{log_btn} out of range 0..{len(order) - 1} "
                              f"for {fix_name!r}")
                continue
            phys_btn = order[log_btn]
            tgt = Target(slave=slave, addr=address)
            existing = bindings[(fix_id, phys_btn)]
            if any(t.slave == tgt.slave and t.addr == tgt.addr for t in existing):
                continue  # duplicate target, silently dedup
            existing.append(tgt)

    for (fix_id, btn), targets in bindings.items():
        if len(targets) > MAX_TARGETS_PER_BTN:
            errors.append(f"Fixture {fix_id} button {btn}: {len(targets)} "
                          f"targets exceeds limit {MAX_TARGETS_PER_BTN}")

    return dict(bindings), errors



# ── bindings.txt emit (replaces sync_bindings.py's serial / replace_mappings) ──

def emit_bindings_txt(bindings: dict) -> str:
    """Render per-(fixture, physical_button) bindings into the positional
    bindings.txt main.lua reads: line index = fixture*8 + button, each line a
    comma-separated list of device*256+address, blank for an unbound button.

    Always emits the full MAX_FIXTURES*MAX_BUTTONS (24*8 = 192) lines so the
    positional read stays aligned — exactly what main.lua's bindings.write()
    produces (for addr=0,#bindings-1)."""
    n = MAX_FIXTURES * MAX_BUTTONS                  # 192 lines, indices 0..191
    lines: list[str] = []
    for idx in range(n):
        fix, btn = divmod(idx, MAX_BUTTONS)
        targets = bindings.get((fix, btn), [])
        encoded = [t.slave * 256 + t.addr for t in targets]
        lines.append(",".join(str(v) for v in encoded))
    return "\n".join(lines) + "\n"


def main() -> int:
    here = Path(__file__).resolve().parent
    default_wb = here.parent / "Switch wiring.xlsx"   # controller/Switch wiring.xlsx

    ap = argparse.ArgumentParser(
        description="Generate a Lua-era /bindings.txt from Switch wiring.xlsx")
    ap.add_argument("workbook", nargs="?", type=Path, default=default_wb,
                    help=f"path to the workbook (default: {default_wb})")
    ap.add_argument("-o", "--output", type=Path, default=Path("bindings.txt"),
                    help="output file (default: ./bindings.txt)")
    ap.add_argument("--stdout", action="store_true",
                    help="write to stdout instead of a file")
    args = ap.parse_args()

    if not args.workbook.exists():
        print(f"workbook not found: {args.workbook}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    remap, errs1 = load_remap(wb)
    bindings, errs2 = load_bindings(wb, remap)

    all_errs = errs1 + errs2
    if all_errs:
        print("Validation errors:", file=sys.stderr)
        for e in all_errs:
            print(f"  ! {e}", file=sys.stderr)
        return 1

    text = emit_bindings_txt(bindings)
    nbtn = sum(1 for v in bindings.values() if v)
    ntgt = sum(len(v) for v in bindings.values())

    if args.stdout:
        sys.stdout.write(text)
    else:
        args.output.write_text(text)
        print(f"wrote {args.output}: {MAX_FIXTURES * MAX_BUTTONS} lines, "
              f"{nbtn} bound buttons, {ntgt} targets", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
